"""Prepare the 2026-10-30 Spanish advisor update from the approved 2026-10-29 deck."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Presentacion_MLSMFP_29-10_avances_auditoria_transferencia.pptx"
OUTPUT = ROOT / "Presentacion_MLSMFP_30-10_avance_chico_norte_y_modelos.pptx"

NAVY = "0E2A47"
BLUE = "355A84"
GOLD = "D6A84F"
GRAY = "5B6977"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_text(shape, value: str, *, size: float | None = None,
             color: str | None = None, bold: bool | None = None,
             align=None) -> None:
    tf = shape.text_frame
    old_p = tf.paragraphs[0] if tf.paragraphs else None
    old_run = old_p.runs[0] if old_p and old_p.runs else None
    old_name = old_run.font.name if old_run else "Montserrat"
    old_size = old_run.font.size.pt if old_run and old_run.font.size else 10
    old_bold = old_run.font.bold if old_run else False
    old_color = None
    if old_run:
        try:
            old_color = old_run.font.color.rgb
        except (AttributeError, TypeError):
            pass
    old_align = old_p.alignment if old_p else PP_ALIGN.LEFT
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align if align is not None else old_align
    p.space_after = Pt(1)
    run = p.add_run()
    run.text = value
    run.font.name = old_name or "Montserrat"
    run.font.size = Pt(size if size is not None else old_size)
    run.font.bold = bold if bold is not None else old_bold
    if color:
        run.font.color.rgb = rgb(color)
    elif old_color:
        run.font.color.rgb = old_color
    tf.word_wrap = True


def set_at(prs, slide_no: int, shape_no: int, value: str, **kwargs) -> None:
    set_text(prs.slides[slide_no - 1].shapes[shape_no], value, **kwargs)


def update_footer_dates(prs) -> None:
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and "29-10-2026" in shape.text:
                set_text(shape, shape.text.replace("29-10-2026", "30-10-2026"))


def comparison_chart() -> BytesIO:
    fig, ax = plt.subplots(figsize=(7.2, 3.3), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    names = ["XGBoost · trial 13", "MLP · trial 48"]
    values = [73.7821, 86.4202]
    bars = ax.barh(names, values, color=[f"#{BLUE}", f"#{GOLD}"], height=0.5)
    ax.invert_yaxis()
    ax.set_xlim(0, 100)
    ax.set_xlabel("Mejor objetivo interno · macro_network_rmse (m³)", fontsize=9)
    ax.set_title("Búsqueda interna en Chico Sur · menor es mejor",
                 fontsize=11, color=f"#{NAVY}", weight="bold", loc="left")
    ax.grid(axis="x", color="#E5EAF0", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.spines["bottom"].set_color("#D7E0E8")
    ax.tick_params(axis="y", length=0, labelsize=9)
    ax.tick_params(axis="x", labelsize=8, colors=f"#{GRAY}")
    for bar, value in zip(bars, values):
        ax.text(value + 1.2, bar.get_y() + bar.get_height() / 2,
                f"{value:.2f} m³", va="center", fontsize=9,
                color=f"#{NAVY}", weight="bold")
    fig.tight_layout(pad=1.2)
    image = BytesIO()
    fig.savefig(image, format="png", bbox_inches="tight", facecolor="white")
    plt.close(fig)
    image.seek(0)
    return image


def replace_model_chart(prs) -> None:
    slide = prs.slides[8]
    old_image = next(shape for shape in slide.shapes if shape.shape_type == 13)
    left, top, width, height = old_image.left, old_image.top, old_image.width, old_image.height
    old_image._element.getparent().remove(old_image._element)
    slide.shapes.add_picture(comparison_chart(), left, top, width=width, height=height)


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    prs = Presentation(SOURCE)
    if len(prs.slides) != 13:
        raise ValueError(f"Unexpected source deck slide count: {len(prs.slides)}")

    update_footer_dates(prs)
    set_at(prs, 1, 3, "30 de octubre de 2026")

    # Executive summary.
    set_at(prs, 2, 1, "Chico Sur lista; Chico Norte incorporada como prueba externa")
    set_at(prs, 2, 2, "Corte al 30 de octubre · candidatos A ajustados; transferencia B aún sin medir.")
    set_at(prs, 2, 6, "CHICO SUR · ENTRENAMIENTO")
    set_at(prs, 2, 7, "138.560 filas")
    set_at(prs, 2, 8, "866 de 1.000 simulaciones elegibles con el QA provisional.", size=9.5)
    set_at(prs, 2, 10, "CANDIDATOS · RMSE INTERNO")
    set_at(prs, 2, 11, "73,78 / 86,42 m³", size=17.5)
    set_at(prs, 2, 12, "XGBoost / MLP · 50 ensayos cada uno; selección solo con Chico Sur.", size=9.0)
    set_at(prs, 2, 14, "CHICO NORTE · CAMPAÑA SWMM")
    set_at(prs, 2, 15, "975 / 975")
    set_at(prs, 2, 16, "787 elegibles; 188 excluidas. Aún no hay score de transferencia.", size=9.1)
    set_at(prs, 2, 18, "MLP")
    set_at(prs, 2, 19, "Búsqueda 50/50 terminada y ensemble final de cinco semillas (42–46) ajustado.", size=10.0)
    set_at(prs, 2, 20, "Chico Norte es la prueba externa. Para medir transferencia: congelar dominio, registrar candidatos reales y sellar predicciones.", size=9.2)

    # Keep the existing A quality figures and clarify their provisional status.
    set_at(prs, 3, 2, "QA hidráulico aplicado a las simulaciones reconstruidas de Chico Sur; criterios aún provisionales.")
    set_at(prs, 3, 33, "134 excluidas: 100 no drenadas y 35 sobre 5% de continuidad; un caso comparte motivos. El umbral y las tolerancias siguen pendientes de justificación hidráulica.", size=8.7)

    # Clarify the actual status of the held-out network.
    set_at(prs, 4, 2, "Chico Sur selecciona los hiperparámetros; Chico Norte se conserva como prueba externa independiente.")
    set_at(prs, 4, 16, "975 corridas staged  Etiquetas y score real aún cerrados", size=9.0)
    set_at(prs, 4, 21, "Los manifiestos enlazan INP, hidrogramas, RPT/OUT, QA y variables. La ruta de scoring está probada con fixtures sintéticos; falta congelar la compatibilidad estructural antes de registrar y predecir en B.", size=9.1)
    set_at(prs, 4, 22, "Una red externa mide transferencia a Chico Norte; no estima desempeño para una población de redes.", size=9.0)

    # Synthetic campaign is now complete, not merely planned.
    set_at(prs, 6, 6, "ESCENARIOS EJECUTADOS")
    set_at(prs, 6, 7, "975")
    set_at(prs, 6, 8, "39 formas × 25 factores; resultados SWMM completos y filtrados por QA.", size=9.3)
    set_at(prs, 6, 9, "QA final: 787 elegibles · 188 excluidas.", size=8.7)
    set_at(prs, 6, 12, "0,2–5,0")
    set_at(prs, 6, 15, "Entrada sintética de caudal por nodo; no es lluvia observada ni transformación lluvia–escorrentía. Se aplicó el redondeo aprobado al minuto más cercano.", size=8.7)

    # Existing spatial comparisons are historical examples, not current candidates.
    set_at(prs, 7, 2, "Mapa SWMM y modelo promovido anteriormente · mismo escenario y escala; comparación retrospectiva, no validación del candidato actual.", size=9.0)
    set_at(prs, 7, 9, "Diferencia del volumen total en este ejemplo: −240 m³. XGBoost anterior (promovido el 06-oct); no es paridad OOF ni prueba externa.", size=8.2)
    set_at(prs, 8, 2, "Mapa SWMM y modelo promovido anteriormente · mismo escenario y escala; comparación retrospectiva, no validación del candidato actual.", size=9.0)
    set_at(prs, 8, 9, "Diferencia del volumen total en este ejemplo: +5.221 m³. XGBoost anterior (promovido el 06-oct); no es paridad OOF ni prueba externa.", size=8.2)

    # Compare current A-only model-selection outcomes on the same internal objective.
    set_at(prs, 9, 1, "Selección interna de XGBoost y MLP")
    set_at(prs, 9, 2, "50 ensayos cada uno · GroupKFold(5) por forma · Chico Sur como única red de selección.", size=9.2)
    set_at(prs, 9, 7, "XGBOOST · MEJOR ENSAYO")
    set_at(prs, 9, 8, "73,782 m³", size=20)
    set_at(prs, 9, 9, "Trial 13 · menor macro_network_rmse.", size=8.7)
    set_at(prs, 9, 11, "MLP · MEJOR ENSAYO")
    set_at(prs, 9, 12, "86,420 m³", size=20)
    set_at(prs, 9, 13, "Trial 48 · mismo objetivo y protocolo.", size=8.7)
    set_at(prs, 9, 15, "Ensemble MLP de cinco semillas ajustado (42–46).", size=9.0)
    set_at(prs, 9, 16, "Mejores objetivos internos; no miden transferencia a B ni generalización entre redes.", size=8.7)
    replace_model_chart(prs)

    # Latest verified campaign and model status.
    set_at(prs, 10, 1, "Candidatos ajustados; Chico Norte simulada y preparada")
    set_at(prs, 10, 2, "Corte verificado al 30 de octubre · conteos de B describen QA, no desempeño del modelo.", size=9.2)
    set_at(prs, 10, 6, "XGBOOST · CANDIDATO A LISTO")
    set_at(prs, 10, 7, "50 / 50 ensayos")
    set_at(prs, 10, 8, "Seleccionado y ajustado con Chico Sur; artefactos y procedencia verificados.", size=9.5)
    set_at(prs, 10, 10, "MLP · ENSEMBLE LISTO")
    set_at(prs, 10, 11, "50 / 50 + 5 semillas")
    set_at(prs, 10, 12, "Trials terminados; ensemble 42–46 recargado y verificado. La selección interna sigue siendo exploratoria.", size=9.1)
    set_at(prs, 10, 13, "CHICO NORTE · QA")
    set_at(prs, 10, 14, "787 elegibles · 188 excluidas · 0 pendientes")
    set_at(prs, 10, 17, "3.901 artefactos enlazados con hashes verificados. QA de continuidad (5%) y drenaje continúa provisional.", size=8.8)

    # Transfer gate and explicit visual inventory.
    set_at(prs, 11, 1, "Transferencia a Chico Norte: paso actual y visuales disponibles")
    set_at(prs, 11, 2, "Las pruebas sintéticas del scorer no son un resultado real de transferencia.", size=9.2)
    set_at(prs, 11, 7, "Congelar dominio")
    set_at(prs, 11, 8, "Aprobar compatibilidad estructural común de A y B.", size=8.8)
    set_at(prs, 11, 11, "Registrar candidatos")
    set_at(prs, 11, 12, "Enlazar modelos, fuentes y variables reales por hashes.", size=8.7)
    set_at(prs, 11, 15, "Predecir y sellar")
    set_at(prs, 11, 16, "Emitir predicciones en B antes de abrir sus etiquetas.", size=8.7)
    set_at(prs, 11, 19, "Evaluar y mapear")
    set_at(prs, 11, 20, "Calcular métricas y producir paridad/mapas del candidato.", size=8.5)
    set_at(prs, 11, 21, "Ya existen: topología de Chico Norte y mapas SWMM de inundación en Chico Sur; las comparaciones mostradas usan un XGBoost anterior. Aún no existen mapas de predicción ni paridad para el candidato nuevo en Chico Norte.", size=9.0)
    set_at(prs, 11, 22, "El scoring real sigue cerrado: falta aprobar la política de dominio y sellar predicciones fuente-enlazadas.", size=8.8)

    # Conclusions and next steps.
    set_at(prs, 12, 1, "Avance logrado y decisiones para completar la transferencia")
    set_at(prs, 12, 2, "Chico Norte ya fue simulada y preparada como prueba externa; todavía no se usa para entrenar.", size=9.4)
    set_at(prs, 12, 7, "• Chico Sur: 866 elegibles; 138.560 filas.", size=8.8)
    set_at(prs, 12, 8, "• XGBoost y MLP: 50/50 trials; ensemble MLP listo.", size=8.8)
    set_at(prs, 12, 9, "• Chico Norte: 975; 787 elegibles; 3.901 hashes verificados.", size=8.5)
    set_at(prs, 12, 12, "• Congelar dominio A→B; QA sigue provisional.", size=8.5)
    set_at(prs, 12, 13, "• Registrar candidatos y sellar predicciones reales.", size=8.5)
    set_at(prs, 12, 14, "• Puntuar B, crear paridad/mapas y probar una tercera red.", size=8.5)
    set_at(prs, 12, 15, "Cautela hidráulica: el umbral de continuidad del 5% y las tolerancias de drenaje siguen provisionales; presentar como filtro operativo, no estándar validado.", size=8.8)
    set_at(prs, 12, 16, "Alcance actual: hidrogramas sintéticos de caudal por nodo. No se simula lluvia–escorrentía.", size=8.8)

    # Compact, current traceability.
    set_at(prs, 13, 1, "Trazabilidad y límites de la evidencia")
    set_at(prs, 13, 2, "Modelos y campaña tienen manifiestos; los objetivos de Chico Norte permanecen cerrados.", size=9.3)
    set_at(prs, 13, 6, "Preparación A")
    set_at(prs, 13, 7, "813323e3ef75b752")
    set_at(prs, 13, 8, "138.560 filas · 866 corridas elegibles · tabular_v4_21", size=9.0)
    set_at(prs, 13, 10, "XGBoost · hashes")
    set_at(prs, 13, 11, "a1a6e504… / 045526a9…")
    set_at(prs, 13, 12, "Clasificador / regresor · candidato A-only", size=8.8)
    set_at(prs, 13, 14, "MLP · hashes")
    set_at(prs, 13, 15, "0bb1045f… / a450253d…")
    set_at(prs, 13, 16, "Clasificador / regresor · ensemble 42–46", size=8.8)
    set_at(prs, 13, 18, "Chico Norte")
    set_at(prs, 13, 19, "975 · 787 QA aptas")
    set_at(prs, 13, 20, "3.901 artefactos con hashes verificados; 188 excluidas", size=8.8)
    set_at(prs, 13, 22, "Scorer · fixtures")
    set_at(prs, 13, 23, "98 pruebas aprobadas")
    set_at(prs, 13, 24, "Pruebas sintéticas; sin score real de B", size=8.8)
    set_at(prs, 13, 25, "Trazabilidad completa en manifests y bitácora; no se incorporaron etiquetas de B al entrenamiento.", size=8.4)

    prs.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    main()
