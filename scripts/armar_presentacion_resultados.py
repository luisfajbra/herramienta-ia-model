"""Agrega 6 slides de resultados (Fase C2 + benchmark) al final del PPT
existente, duplicando el estilo visual de las slides 17-29 (titulo + subtitulo
+ tabla/imagenes + pie de pagina + numero), sin tocar ninguna slide original.

Nunca sobreescribe el .pptx original -- siempre guarda un archivo nuevo.

Uso:
    venv/Scripts/python.exe scripts/armar_presentacion_resultados.py
"""

from __future__ import annotations

import copy
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation
from pptx.util import Emu, Pt

ORIGINAL = Path("Presentacion_MLSMFP_15-09_top5_resultados.pptx")
OUTPUT = Path("Presentacion_MLSMFP_05-10_resultados_fase_c2.pptx")
PARIDAD_DIR = Path("outputs/bench/reports/paridad")

FOOTER_TEXT = "MLSMFP  |  Resultados Fase C2 · benchmark 5 familias  |  05 octubre 2026"

FAMILY_LABELS = {
    "xgboost": "XGBoost",
    "random_forest": "Random Forest",
    "linear": "Lineal",
    "mlp": "MLP",
}


def _copy_slide_images(source_slide, dest_slide) -> dict[str, str]:
    """Copia las relaciones de imagen del slide origen al destino.

    Devuelve un mapa {rId_original: rId_nuevo} para poder re-escribir los
    r:embed de las copias de <pic> que se peguen despues.
    """
    rid_map = {}
    for rel_id, rel in source_slide.part.rels.items():
        if "image" in rel.reltype:
            new_rid = dest_slide.part.rels._next_rId if hasattr(dest_slide.part.rels, "_next_rId") else None
            new_rel = dest_slide.part.rels.add_relationship(
                rel.reltype, rel.target_part, None, is_external=False
            )
            rid_map[rel_id] = new_rel
    return rid_map


def duplicate_slide(prs: Presentation, index: int):
    """Duplica prs.slides[index] al final de prs, copiando shapes e imagenes."""
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)

    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)

    rid_map = {}
    for rel_id, rel in source.part.rels.items():
        if "image" in rel.reltype:
            new_rel_id = dest.part.get_or_add_image_part(rel.target_part.blob)[1]
            rid_map[rel_id] = new_rel_id

    from pptx.oxml.ns import qn

    for shp in source.shapes:
        el = copy.deepcopy(shp._element)
        for blip in el.iter(qn("a:blip")):
            old_rid = blip.get(qn("r:embed"))
            if old_rid in rid_map:
                blip.set(qn("r:embed"), rid_map[old_rid])
        dest.shapes._spTree.append(el)

    return dest


def _set_text(slide, contains: str, new_text: str) -> None:
    for shape in slide.shapes:
        if shape.has_text_frame and contains in shape.text_frame.text:
            tf = shape.text_frame
            first_run = tf.paragraphs[0].runs[0]
            for p in list(tf.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            for r in list(tf.paragraphs[0].runs)[1:]:
                r._r.getparent().remove(r._r)
            first_run.text = new_text
            return
    raise ValueError(f"No se encontro texto con {contains!r}")


def _set_page_number(slide, n: int) -> None:
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip().isdigit():
            shape.text_frame.paragraphs[0].runs[0].text = str(n)
            return


def _fill_table(slide, rows: list[list[str]]) -> None:
    """Reescribe una tabla existente in-place con `rows` (incluye header).

    Si `rows` tiene mas filas que la tabla original, agrega filas clonando
    el alto/estilo de la ultima fila de datos.
    """
    for shape in slide.shapes:
        if shape.has_table:
            tbl = shape.table
            existing_rows = len(tbl.rows)
            needed = len(rows)
            if needed > existing_rows:
                last_tr = tbl._tbl.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}tr")[-1]
                for _ in range(needed - existing_rows):
                    new_tr = copy.deepcopy(last_tr)
                    last_tr.addnext(new_tr)
                    last_tr = new_tr
            for r_idx, row_values in enumerate(rows):
                for c_idx, value in enumerate(row_values):
                    if c_idx >= len(tbl.rows[r_idx].cells):
                        continue
                    cell = tbl.rows[r_idx].cells[c_idx]
                    tf = cell.text_frame
                    if tf.paragraphs[0].runs:
                        base_run = tf.paragraphs[0].runs[0]
                        for p in list(tf.paragraphs)[1:]:
                            p._p.getparent().remove(p._p)
                        for run in list(tf.paragraphs[0].runs)[1:]:
                            run._r.getparent().remove(run._r)
                        base_run.text = value
                    else:
                        tf.text = value
            return
    raise ValueError("No se encontro tabla en el slide")


def _clear_table(slide) -> None:
    """Elimina el shape de tabla de un slide (para reusarlo como lienzo de imagenes)."""
    for shape in list(slide.shapes):
        if shape.has_table:
            shape._element.getparent().remove(shape._element)
            return


def _add_picture_safe(slide, image_path: Path, left, top, width=None, height=None):
    if not image_path.exists():
        print(f"  AVISO: no existe {image_path}, se omite")
        return None
    return slide.shapes.add_picture(str(image_path), Emu(left), Emu(top), Emu(width) if width else None, Emu(height) if height else None)


def main() -> None:
    if not ORIGINAL.exists():
        raise SystemExit(f"No existe {ORIGINAL}")
    prs = Presentation(str(ORIGINAL))
    base_index = len(prs.slides)  # las slides nuevas empiezan despues de esta cantidad
    donor_table_slide = 16  # slide 17 (0-indexed), title+subtitle+tabla+footer+numero

    # --- Slide: Resumen del dataset real -----------------------------------
    s = duplicate_slide(prs, donor_table_slide)
    _set_text(s, "¿Vale la pena cambiar XGBoost?", "Primera corrida real de punta a punta (2026-10-05)")
    _set_text(s, "Objetivo: inundación", "39 formas × 25 factores + base = 1000 corridas SWMM, 160 nodos, 160,000 filas")
    _fill_table(s, [
        ["Métrica", "Valor", "Nota"],
        ["Corridas SWMM", "1000", "39 formas × 25 factores + base"],
        ["Filas en dataset", "160,000", "1000 corridas × 160 nodos"],
        ["Escenarios sin drenar", "0", "Fase C2: reintento con extensión nunca se activó"],
        ["Errores / crashes", "0", "Pipeline completo, 0 fallos"],
    ])
    _set_text(s, "MLSMFP  |", FOOTER_TEXT)
    _set_page_number(s, base_index + 1)

    # --- Slide: Exclusión por continuidad (Fase C2) ------------------------
    s = duplicate_slide(prs, donor_table_slide)
    _set_text(s, "¿Vale la pena cambiar XGBoost?", "Exclusión por continuidad de routing — Fase C2")
    _set_text(s, "Objetivo: inundación", "37 de 1000 corridas (3.7%) superan el 5% de error de continuidad SWMM")
    _fill_table(s, [
        ["Escenario", "Factor", "Error de continuidad"],
        ["shape_ultra_short_impulse", "0.2", "12.38% (peor caso)"],
        ["rare_extreme", "0.2", "11.72%"],
        ["common_bogota", "0.2", "10.57%"],
        ["base (forma base)", "0.2", "5.17% — confirma el fix de formas base"],
    ])
    _set_text(s, "MLSMFP  |", FOOTER_TEXT)
    _set_page_number(s, base_index + 2)

    # --- Slide: Ranking de las 5 familias -----------------------------------
    s = duplicate_slide(prs, donor_table_slide)
    _set_text(s, "¿Vale la pena cambiar XGBoost?", "Benchmark justo — 5 familias, protocolo LeaveOneShapeOut")
    _set_text(s, "Objetivo: inundación", "154,080 filas evaluadas (tras excluir continuidad) · 46,225 inundadas")
    _fill_table(s, [
        ["Familia", "RMSE volumen", "F1 clasificador", "NSE regresor"],
        ["XGBoost (ganador)", "78.04", "0.979", "0.815"],
        ["MLP", "99.61", "0.977", "0.631"],
        ["Random Forest", "185.46", "0.890", "0.277"],
        ["Lineal", "201.18", "0.817", "0.071"],
        ["SVM", "pendiente — excluido por escala", "—", "—"],
    ])
    _set_text(s, "MLSMFP  |", FOOTER_TEXT)
    _set_page_number(s, base_index + 3)

    # --- Slide: Mapas de paridad por familia (grid 2x2) --------------------
    s = duplicate_slide(prs, donor_table_slide)
    _set_text(s, "¿Vale la pena cambiar XGBoost?", "Paridad SWMM vs. predicción, por familia (OOF)")
    _set_text(s, "Objetivo: inundación", "Mismas predicciones out-of-fold que el ranking — no es un ajuste optimista")
    _clear_table(s)
    slide_w, slide_h = 9144000, 5143500
    margin = 300000
    top0 = 1000000
    cell_w = (slide_w - 2 * margin) // 2
    cell_h = (slide_h - top0 - margin) // 2
    positions = [
        (margin, top0), (margin + cell_w, top0),
        (margin, top0 + cell_h), (margin + cell_w, top0 + cell_h),
    ]
    for (family, label), (left, top) in zip(FAMILY_LABELS.items(), positions):
        img = PARIDAD_DIR / family / "parity_nodes_linear_todas_las_corridas.png"
        pic = _add_picture_safe(s, img, left + 50000, top + 50000, height=cell_h - 350000)
        tb = s.shapes.add_textbox(Emu(left + 50000), Emu(top), Emu(cell_w - 100000), Emu(300000))
        tb.text_frame.text = label
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        tb.text_frame.paragraphs[0].runs[0].font.bold = True
    _set_text(s, "MLSMFP  |", FOOTER_TEXT)
    _set_page_number(s, base_index + 4)

    # --- Slide: Optuna sobre XGBoost ----------------------------------------
    s = duplicate_slide(prs, donor_table_slide)
    _set_text(s, "¿Vale la pena cambiar XGBoost?", "Optuna sobre XGBoost (familia ganadora) — estimación exploratoria")
    _set_text(s, "Objetivo: inundación", "50 trials · GroupKFold(5) interno por shape_id · NO es confirmatorio (spec §12)")
    _fill_table(s, [
        ["Hiperparámetro", "Default", "Mejor trial (#47/50)"],
        ["Clasificador max_depth / lr / n_est.", "6 / 0.05 / 200", "5 / 0.136 / 313"],
        ["Regresor max_depth / lr / n_est.", "6 / 0.05 / 200", "8 / 0.199 / 421"],
        ["Objetivo (RMSE, fold interno)", "—", "108.84"],
        ["Status", "—", "estimacion_exploratoria_condicionada"],
    ])
    _set_text(s, "MLSMFP  |", FOOTER_TEXT)
    _set_page_number(s, base_index + 5)

    # --- Slide: Próximos pasos ----------------------------------------------
    s = duplicate_slide(prs, donor_table_slide)
    _set_text(s, "¿Vale la pena cambiar XGBoost?", "Próximos pasos")
    _set_text(s, "Objetivo: inundación", "Lo que queda en la lista de trabajo después de esta corrida")
    _fill_table(s, [
        ["#", "Pendiente", "Origen"],
        ["1", "Configurar SVM para esta escala y sumarlo al ranking", "Benchmark"],
        ["2", "Entrenar sobre tabular_v4_21 (features tributarias) y conectar ScenarioPredictor", "Fase B2"],
        ["3", "config_7shapes.yaml desactualizado", "Auditoría Fase A"],
        ["4", "Decisión de integración (merge a la rama real)", "Pospuesta"],
    ])
    _set_text(s, "MLSMFP  |", FOOTER_TEXT)
    _set_page_number(s, base_index + 6)

    prs.save(str(OUTPUT))
    print(f"Guardado: {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
