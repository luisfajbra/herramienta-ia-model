"""Build the Spanish 2026-10-28 weekly update in the 2026-10-20 deck style."""

from __future__ import annotations

import io
import hashlib
import json
import sys
from pathlib import Path

import joblib
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from swmm_resilience.visualization._inp_parser import parse_conduits, parse_coordinates
from swmm_resilience.simulation.hydrograph_shapes import load_shape
from swmm_resilience.simulation.swmm_api_io import get_node_inflow_profiles, load_inp

TEMPLATE = ROOT / ".claude/worktrees/optuna-tuning-familia-ganadora/Presentacion_MLSMFP_20-10_avances_auditoria_svm.pptx"
EXPERIMENT = ROOT / "outputs/experiments/frozen_pair_rerun_20261028/20261028T194035Z"
PREP_ID = "813323e3ef75b752"
STUDY_PATH = EXPERIMENT / "A/model_selection_Aonly_813323e3ef75b752/studies/xgboost/813323e3ef75b752_20261028T224207Z/study.json"
QUALITY_PATH = EXPERIMENT / f"A/prepared_a_transfer_v1/{PREP_ID}/quality_report.json"
PREREG_PATH = EXPERIMENT / "A/model_selection_Aonly_813323e3ef75b752/selection_preregistration_v2.json"
B_PROGRESS_PATH = EXPERIMENT / "B/batch_progress.json"
PREPARED_A = EXPERIMENT / f"A/prepared_a_transfer_v1/{PREP_ID}"
A_INP = ROOT / "data/networks/chico_hydro-qx1/SWMM - Chico (PVC) Prueba 1 - Qx1.00.inp"
B_INP = ROOT / "data/networks/chico_norte_93_qx1/SWMM - Chico Norte 93% - Qx1.00.inp"
VISUALS = EXPERIMENT / "presentation_visuals_modelo_hidrograma"
PROMOTED_MODELS = ROOT / "outputs/models"
OUTPUT = ROOT / "Presentacion_MLSMFP_28-10_avances_auditoria_transferencia_comparacion_modelo_hidrograma.pptx"

# Palette sampled from the 20-10 presentation: navy, layered blues and gold.
NAVY = RGBColor(14, 42, 71)          # 0E2A47
BLUE = RGBColor(53, 90, 132)         # 355A84
MID_BLUE = RGBColor(94, 140, 190)    # 5E8CBE
LIGHT_BLUE = RGBColor(139, 197, 250) # 8BC5FA
GOLD = RGBColor(241, 203, 131)       # F1CB83
PALE_BLUE = RGBColor(233, 242, 250)
PALE_GOLD = RGBColor(251, 245, 232)
PALE_GRAY = RGBColor(242, 245, 248)
GRAY = RGBColor(91, 105, 119)
BLACK = RGBColor(45, 48, 52)
WHITE = RGBColor(255, 255, 255)


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def add_text(slide, x, y, w, h, value, *, size=11, color=BLACK,
             bold=False, font="Montserrat", align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.MIDDLE, margin=0.06):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def add_box(slide, x, y, w, h, fill=PALE_BLUE, *, square=False):
    kind = MSO_SHAPE.RECTANGLE if square else MSO_SHAPE.ROUNDED_RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    shape.line.width = Pt(0.6)
    if not square:
        try:
            shape.adjustments[0] = 0.08
        except (IndexError, ValueError):
            pass
    return shape


def add_card(slide, x, y, w, h, label, value, detail, *, fill=PALE_BLUE,
             accent=BLUE, value_size=21):
    add_box(slide, x, y, w, h, fill)
    add_text(slide, x + 0.14, y + 0.08, w - 0.28, 0.24,
             label.upper(), size=8.5, color=accent, bold=True)
    add_text(slide, x + 0.14, y + 0.36, w - 0.28, 0.43,
             value, size=value_size, color=NAVY, bold=True)
    add_text(slide, x + 0.14, y + 0.82, w - 0.28, h - 0.88,
             detail, size=9.1, color=BLACK, valign=MSO_ANCHOR.TOP)


def add_header(slide, title, subtitle, page):
    add_text(slide, 0.44, 0.17, 9.12, 0.48, title, size=21.5,
             color=NAVY, bold=True, font="Libre Baskerville")
    add_text(slide, 0.47, 0.67, 9.05, 0.35, subtitle,
             size=9.7, color=GRAY)
    add_text(slide, 0.34, 5.32, 8.6, 0.17,
             "MLSMFP | Avance: auditoría, ajuste y transferencia entre redes | 28-10-2026",
             size=7.4, color=GRAY)
    add_text(slide, 9.1, 5.25, 0.55, 0.25, str(page), size=11,
             color=GRAY, align=PP_ALIGN.CENTER)


def add_slide(prs, title, subtitle, page):
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    add_header(slide, title, subtitle, page)
    return slide


def fmt_int(n):
    return f"{int(n):,}".replace(",", ".")


def fmt_decimal(n, digits=1):
    return f"{float(n):.{digits}f}".replace(".", ",")


def chart_bytes(study):
    values = [t.get("value") for t in study["trials"]]
    best_so_far = []
    best = float("inf")
    for value in values:
        if value is not None:
            best = min(best, float(value))
        best_so_far.append(best)
    best_value = float(study["best_value"])
    best_idx = next(i for i, v in enumerate(values) if v == best_value)
    x = list(range(1, len(values) + 1))

    fig, ax = plt.subplots(figsize=(8.4, 4.0), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.plot(x, values, color="#8BC5FA", marker="o", markersize=3,
            linewidth=1.1, label="RMSE de cada ensayo")
    ax.plot(x, best_so_far, color="#355A84", linewidth=2.6,
            label="Mejor RMSE acumulado")
    ax.scatter([best_idx + 1], [best_value], s=55, color="#F1CB83", zorder=5)
    ax.annotate(f"Mejor: {best_value:.2f} m\u00b3\nEnsayo Optuna {best_idx}",
                (best_idx + 1, best_value), xytext=(22, 22),
                textcoords="offset points", fontsize=9,
                arrowprops={"arrowstyle": "->", "color": "#5B6977"})
    ax.set_xlabel("Ensayo Optuna (orden de ejecuci\u00f3n)", fontsize=10)
    ax.set_ylabel("RMSE interno de GroupKFold(5) (m\u00b3)", fontsize=10)
    ax.grid(axis="y", color="#E5EAF0", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_xlim(1, len(values))
    ax.legend(frameon=False, loc="upper right", fontsize=8.5)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


def draw_topology(inp_path: Path, out_path: Path):
    coords = parse_coordinates(inp_path)
    conduits = parse_conduits(inp_path)
    fig, ax = plt.subplots(figsize=(8, 6), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    for _link, frm, to in conduits:
        if frm in coords and to in coords:
            x0, y0 = coords[frm]
            x1, y1 = coords[to]
            ax.plot([x0, x1], [y0, y1], color="#355A84", linewidth=0.62, alpha=0.82, zorder=1)
    xs = [xy[0] for xy in coords.values()]
    ys = [xy[1] for xy in coords.values()]
    ax.scatter(xs, ys, s=8, color="#8BC5FA", edgecolors="#0E2A47", linewidths=0.2, zorder=2)
    ax.set_title("Red B · Chico Norte 93% · topología del INP", color="#0E2A47", fontsize=12, weight="bold")
    ax.set_aspect("equal")
    ax.axis("off")
    ax.text(0.02, 0.02, f"{len(coords)} nodos con coordenadas · {len(conduits)} conducciones",
            transform=ax.transAxes, fontsize=8.5, color="#5B6977")
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return out_path, len(coords), len(conduits)


def draw_flood_map(inp_path: Path, rows: pd.DataFrame, title: str, out_path: Path,
                   *, vmax: float | None = None):
    coords = parse_coordinates(inp_path)
    conduits = parse_conduits(inp_path)
    df = rows.copy()
    df["node_id"] = df["node_id"].astype(str)
    df["vol_inundacion_m3"] = pd.to_numeric(df["vol_inundacion_m3"], errors="raise").clip(lower=0)
    df = df[df["node_id"].isin(coords)].drop_duplicates("node_id")
    df["x"] = df["node_id"].map(lambda node: coords[node][0])
    df["y"] = df["node_id"].map(lambda node: coords[node][1])
    vmax = max(float(vmax or df["vol_inundacion_m3"].max()), 1.0)
    fig, ax = plt.subplots(figsize=(7.0, 5.2), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    for _link, frm, to in conduits:
        if frm in coords and to in coords:
            ax.plot([coords[frm][0], coords[to][0]], [coords[frm][1], coords[to][1]],
                    color="#B8C4D0", linewidth=0.62, zorder=1)
    dry = df[df["vol_inundacion_m3"] <= 0]
    wet = df[df["vol_inundacion_m3"] > 0]
    ax.scatter(dry["x"], dry["y"], s=13, color="#DCE6EF", edgecolors="#8AA0B4",
               linewidths=0.25, label="Sin inundación", zorder=2)
    if not wet.empty:
        sizes = 22 + 120 * wet["vol_inundacion_m3"] / vmax
        points = ax.scatter(wet["x"], wet["y"], s=sizes, c=wet["vol_inundacion_m3"],
                            cmap="YlOrRd", vmin=0, vmax=vmax, edgecolors="white",
                            linewidths=0.35, label="Con inundación", zorder=3)
        cb = fig.colorbar(points, ax=ax, fraction=0.035, pad=0.025)
        cb.set_label("Volumen de inundación (m³)", fontsize=8.5)
    top = wet.nlargest(4, "vol_inundacion_m3")
    for _, row in top.iterrows():
        ax.annotate(f"{row['node_id']} · {row['vol_inundacion_m3']:,.0f} m³",
                    (row["x"], row["y"]), xytext=(5, 5), textcoords="offset points",
                    fontsize=6.7, color="#0E2A47",
                    bbox={"boxstyle": "round,pad=0.15", "fc": "white", "ec": "#D7E0E8", "alpha": 0.9})
    ax.set_title(title, fontsize=10.5, color="#0E2A47", weight="bold", pad=8)
    ax.set_xlabel("Coordenada X (m)", fontsize=8)
    ax.set_ylabel("Coordenada Y (m)", fontsize=8)
    ax.set_aspect("equal")
    ax.grid(True, linestyle="--", alpha=0.20)
    ax.legend(loc="upper right", framealpha=0.92, fontsize=7.3)
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return out_path


def draw_b_hydrograph(inp_path: Path, out_path: Path):
    inp = load_inp(inp_path)
    profiles = get_node_inflow_profiles(inp)
    peaks = {
        node: max((flow for _, flow in profile["points"]), default=0.0)
        for node, profile in profiles.items()
    }
    node = max(peaks, key=peaks.get)
    base_peak = peaks[node]
    factor = 2.8
    shape_name = "shape_five_pulse_8h5"
    points = load_shape(ROOT / "data/hydrograph_shapes" / f"{shape_name}.csv")
    times = np.asarray([round(t * 60) / 60 for t, _ in points])
    normalized = np.asarray([q for _, q in points])
    flow = normalized * base_peak * factor
    fig, ax = plt.subplots(figsize=(7.8, 3.7), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.plot(times, flow, color="#355A84", linewidth=2.6, marker="o", markersize=3.6)
    ax.fill_between(times, flow, color="#8BC5FA", alpha=0.24)
    ax.set_title(f"Hidrograma sintético de entrada · nodo {node}",
                 color="#0E2A47", fontsize=12, weight="bold")
    ax.set_xlabel("Tiempo desde el inicio del evento (h)", fontsize=9)
    ax.set_ylabel("Caudal de entrada (L/s)", fontsize=9)
    ax.grid(True, color="#E5EAF0", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.text(0.99, 0.96,
            f"{shape_name} · {fmt_decimal(points[-1][0], 1)} h · factor {fmt_decimal(factor, 1)}\n"
            f"pico INP {fmt_decimal(base_peak, 2)} → aplicado {fmt_decimal(flow.max(), 1)} L/s",
            transform=ax.transAxes, ha="right", va="top", fontsize=8.0,
            color="#0E2A47", bbox={"boxstyle": "round,pad=0.35", "fc": "white", "ec": "#D7E0E8"})
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return out_path, {"node": node, "base_peak_lps": base_peak,
                      "scenario_peak_lps": float(flow.max()), "factor": factor,
                      "shape": shape_name, "duration_h": float(points[-1][0]),
                      "inflow_nodes": len(profiles), "positive_inflow_nodes": sum(p > 0 for p in peaks.values())}


def make_visuals():
    VISUALS.mkdir(parents=True, exist_ok=True)
    network_path, n_nodes_b, n_links_b = draw_topology(B_INP, VISUALS / "red_b_chico_norte_topologia.png")
    hydrograph_path, hydrograph_stats = draw_b_hydrograph(B_INP, VISUALS / "red_b_hidrograma_sintetico.png")
    keys = pd.read_parquet(PREPARED_A / "keys.parquet")
    features = pd.read_parquet(PREPARED_A / "features.parquet")
    targets = pd.read_parquet(PREPARED_A / "targets.parquet")
    joined = keys[["run_id", "scenario_key", "shape_id", "factor_mult", "node_id"]].copy()
    joined["vol_inundacion_m3"] = targets["vol_inundacion_m3"].to_numpy()
    joined["inunda"] = targets["inunda"].to_numpy()
    summary = joined.groupby(["run_id", "scenario_key", "shape_id", "factor_mult"], as_index=False).agg(
        total_m3=("vol_inundacion_m3", "sum"), n_flooded=("inunda", "sum"))
    flooded = summary[summary["n_flooded"] > 0].sort_values(["total_m3", "scenario_key"])
    picks = [flooded.iloc[len(flooded) // 2], flooded.iloc[-1]]
    promotion = read_json(PROMOTED_MODELS / "promotion.json")
    if promotion.get("feature_contract_id") != "tabular_v4_21":
        raise ValueError("Previous promoted model is not compatible with tabular_v4_21")
    model_paths = {"classifier": PROMOTED_MODELS / "classifier.joblib",
                   "regressor": PROMOTED_MODELS / "regressor.joblib"}
    for key, model_path in model_paths.items():
        expected = promotion[f"{key}_sha256"]
        actual = hashlib.sha256(model_path.read_bytes()).hexdigest()
        if actual != expected:
            raise ValueError(f"Promoted {key} hash does not match promotion.json")
    classifier = joblib.load(model_paths["classifier"])
    regressor = joblib.load(model_paths["regressor"])
    if list(features.columns) != promotion["ordered_features"]:
        raise ValueError("Prepared A feature columns differ from the previous model contract")
    maps = []
    for label, pick in zip(("mediana", "mayor_volumen"), picks):
        run_id = int(pick["run_id"])
        row_index = keys.index[keys["run_id"].eq(run_id)]
        scenario_rows = joined.loc[row_index].copy()
        scenario_features = features.loc[row_index]
        probability = classifier.predict_proba(scenario_features)[:, 1]
        conditional = np.maximum(np.expm1(regressor.predict(scenario_features)), 0.0)
        predicted_volume = np.where(probability >= 0.5, conditional, 0.0)
        scenario_rows["predicted_volume_m3"] = predicted_volume
        model_rows = scenario_rows.drop(columns=["vol_inundacion_m3", "predicted_volume_m3"]).copy()
        model_rows["vol_inundacion_m3"] = predicted_volume
        total_true = float(scenario_rows["vol_inundacion_m3"].sum())
        total_model = float(predicted_volume.sum())
        common_vmax = max(float(scenario_rows["vol_inundacion_m3"].max()), float(predicted_volume.max()), 1.0)
        scenario_title = f"{pick['shape_id']} · factor {fmt_decimal(pick['factor_mult'], 1)}"
        out_swmm = draw_flood_map(
            A_INP, scenario_rows, f"SWMM · {scenario_title}",
            VISUALS / f"red_a_{label}_swmm.png", vmax=common_vmax)
        out_model = draw_flood_map(
            A_INP, model_rows,
            f"XGBoost previo · {scenario_title}",
            VISUALS / f"red_a_{label}_xgboost_previo.png", vmax=common_vmax)
        stats = pick.to_dict()
        stats["model_total_m3"] = total_model
        stats["total_delta_m3"] = total_model - total_true
        stats["model_flooded_nodes"] = int((predicted_volume >= 1.0).sum())
        maps.append({"swmm": out_swmm, "model": out_model, "stats": stats})
    return {"network": network_path, "b_nodes": n_nodes_b, "b_links": n_links_b,
            "b_hydrograph": hydrograph_path, "b_hydrograph_stats": hydrograph_stats,
            "a_maps": maps, "previous_model_promoted_utc": promotion["promoted_at_utc"]}


def keep_cover_only(prs):
    ids = prs.slides._sldIdLst
    keep = ids[43]
    for item in list(ids):
        if item is not keep:
            rid = item.rId
            ids.remove(item)
            prs.part.drop_rel(rid)


def main():
    for path in (TEMPLATE, STUDY_PATH, QUALITY_PATH, PREREG_PATH, B_PROGRESS_PATH,
                 A_INP, B_INP, PREPARED_A / "keys.parquet", PREPARED_A / "features.parquet",
                 PREPARED_A / "targets.parquet", PROMOTED_MODELS / "promotion.json",
                 PROMOTED_MODELS / "classifier.joblib", PROMOTED_MODELS / "regressor.joblib",
                 ROOT / "data/hydrograph_shapes/shape_five_pulse_8h5.csv"):
        if not path.exists():
            raise FileNotFoundError(path)
    study = read_json(STUDY_PATH)
    quality = read_json(QUALITY_PATH)
    b_progress = read_json(B_PROGRESS_PATH)
    prereg = read_json(PREREG_PATH)
    visuals = make_visuals()

    prs = Presentation(str(TEMPLATE))
    keep_cover_only(prs)
    cover = prs.slides[0]
    for shape in cover.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text.strip()
        if text == "2026-10-20":
            shape.text = "28 de octubre de 2026"
        elif text == "Machine Learning Model for Hydraulic Failure Prediction in Stormwater Systems":
            shape.text = "Modelo de aprendizaje autom\u00e1tico para predecir fallas hidr\u00e1ulicas\nen redes de alcantarillado pluvial"

    # 45. Executive summary: three clear takeaways.
    s = add_slide(prs, "Avance semanal: datos A listos; transferencia en curso",
                  "Corte al 28 de octubre de 2026 · distinguir selecci\u00f3n interna de evaluaci\u00f3n externa.", 45)
    add_card(s, 0.52, 1.36, 2.76, 1.58, "Datos de A", "138.560 filas",
             "866 de 1.000 simulaciones pasaron QA.", fill=PALE_BLUE, accent=BLUE, value_size=18)
    add_card(s, 3.62, 1.36, 2.76, 1.58, "XGBoost", "73,78 m\u00b3",
             "Mejor RMSE interno tras 50 ensayos Optuna en A.", fill=PALE_BLUE, accent=BLUE, value_size=22)
    add_card(s, 6.72, 1.36, 2.76, 1.58, "Prueba externa B", "En curso",
             "A\u00fan no hay m\u00e9tricas de transferencia.", fill=PALE_GOLD, accent=NAVY, value_size=21)
    add_box(s, 0.52, 3.30, 8.96, 0.88, PALE_GRAY)
    add_text(s, 0.76, 3.43, 1.35, 0.34, "MLP", size=16, color=BLUE, bold=True)
    add_text(s, 2.10, 3.40, 7.05, 0.48,
             "B\u00fasqueda de 50 ensayos activa; despu\u00e9s se ajustar\u00e1 el ensemble final de cinco semillas.",
             size=11.2, color=BLACK)
    add_text(s, 0.62, 4.48, 8.85, 0.42,
             "Conclusi\u00f3n de hoy: A es trazable y XGBoost tiene candidato reproducible; la transferencia a B todav\u00eda debe medirse.",
             size=10.6, color=NAVY, bold=True)

    # 46. Dataset A and QA.
    s = add_slide(prs, "Conjunto A: simulaciones filtradas con QA hidr\u00e1ulico",
                  "Criterios aplicados provisionalmente: continuidad, drenaje y cobertura completa de resultados.", 46)
    eligible = int(quality["n_runs"])
    excluded = 1000 - eligible
    cards = [
        (0.52, "Corridas fuente", "1.000", "Simulaciones SWMM"),
        (2.82, "Aptas para entrenar", fmt_int(eligible), "Tras controles de QA"),
        (5.12, "Excluidas", str(excluded), "Sin imputar etiquetas"),
        (7.42, "Formas presentes", "36 / 40", "Cuatro grupos sin corrida apta"),
    ]
    for x, label, value, note in cards:
        add_card(s, x, 1.25, 2.05, 1.22, label, value, note,
                 fill=PALE_GOLD if label == "Excluidas" else PALE_BLUE,
                 accent=NAVY if label == "Excluidas" else BLUE,
                 value_size=20)
    add_text(s, 0.62, 2.77, 5.9, 0.30, "MOTIVOS DE EXCLUSI\u00d3N", size=9.8, color=BLUE, bold=True)
    for i, (label, count) in enumerate([("No drenada", 100), ("Continuidad > 5%", 35)]):
        y = 3.22 + i * 0.64
        add_text(s, 0.65, y, 2.15, 0.30, label, size=10.2, color=BLACK)
        add_box(s, 2.80, y + 0.04, 3.12, 0.20, PALE_GRAY, square=True)
        width = 3.12 * count / 100
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.80), Inches(y + 0.04), Inches(width), Inches(0.20))
        bar.fill.solid(); bar.fill.fore_color.rgb = BLUE if i == 0 else LIGHT_BLUE; bar.line.color.rgb = bar.fill.fore_color.rgb
        add_text(s, 6.02, y - 0.01, 0.48, 0.30, str(count), size=10, color=NAVY, bold=True)
    add_box(s, 7.10, 2.83, 2.35, 1.49, PALE_BLUE)
    add_text(s, 7.27, 2.97, 2.00, 0.31, "138.560 filas", size=15, color=NAVY, bold=True)
    add_text(s, 7.27, 3.36, 2.01, 0.78,
             "160 nodos · 21 variables\n41.497 filas inundadas (29,9%)\nUmbral: 1 m\u00b3",
             size=9.3, color=BLACK, valign=MSO_ANCHOR.TOP)
    add_text(s, 0.62, 4.55, 8.9, 0.36,
             "Los motivos se traslapan en una corrida. El 5% y las tolerancias de drenaje a\u00fan requieren justificaci\u00f3n hidr\u00e1ulica.",
             size=9.2, color=GRAY)

    # 47. Validation boundary.
    s = add_slide(prs, "Dise\u00f1o de evaluaci\u00f3n: B permanece aislada",
                  "La selecci\u00f3n usa A; la red B se reserva para una prueba externa.", 47)
    for x, label, value, detail, fill, accent in [
        (0.52, "Entrenamiento", "Red A · Chico Sur", "866 corridas aptas\nPreparaci\u00f3n congelada", PALE_BLUE, BLUE),
        (3.62, "Selecci\u00f3n interna", "GroupKFold(5)", "Separaci\u00f3n por forma\nOptuna solo ve A", PALE_BLUE, BLUE),
        (6.72, "Prueba externa", "Red B · Chico Norte", "Etiquetas cerradas\nUna red no representa todas", PALE_GOLD, NAVY),
    ]:
        add_card(s, x, 1.35, 2.76, 1.55, label, value, detail, fill=fill, accent=accent, value_size=15)
    add_text(s, 3.19, 1.85, 0.35, 0.42, "→", size=23, color=LIGHT_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, 6.29, 1.85, 0.35, 0.42, "→", size=23, color=LIGHT_BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_box(s, 0.52, 3.20, 8.96, 1.10, PALE_GRAY)
    add_text(s, 0.76, 3.31, 8.4, 0.31, "Control contra fuga de informaci\u00f3n", size=12.5, color=NAVY, bold=True)
    add_text(s, 0.76, 3.68, 8.35, 0.49,
             "Se verifican manifiestos, archivos por corrida, claves y variables. Las predicciones se sellan y reproducen antes de abrir los RPT de B.",
             size=10, color=BLACK)
    add_text(s, 0.62, 4.58, 8.9, 0.30,
             "Alcance: transferencia a una red externa; no prueba generalizaci\u00f3n a una poblaci\u00f3n de redes.",
             size=9.4, color=GRAY)

    # 48. Show B geometry without implying that its hydraulics have been scored.
    s = add_slide(prs, "Red B · Chico Norte: topolog\u00eda del INP", 
                  "Vista de nodos y conducciones; este mapa no contiene resultados de simulaci\u00f3n ni predicciones.", 48)
    s.shapes.add_picture(str(visuals["network"]), Inches(0.35), Inches(1.05), width=Inches(5.00))
    add_card(s, 6.20, 1.40, 3.10, 1.30, "Nodos con coordenadas", fmt_int(visuals["b_nodes"]),
             "Inventario geom\u00e9trico le\u00eddo del INP base.", fill=PALE_BLUE, accent=BLUE, value_size=22)
    add_card(s, 6.20, 2.95, 3.10, 1.30, "Conducciones", fmt_int(visuals["b_links"]),
             "Trazado de enlaces entre nodos seg\u00fan el INP.", fill=PALE_BLUE, accent=BLUE, value_size=22)
    add_box(s, 6.20, 4.48, 3.10, 0.44, PALE_GOLD)
    add_text(s, 6.35, 4.51, 2.80, 0.34,
             "Topolog\u00eda solamente · sin caudales ni inundaci\u00f3n",
             size=8.4, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    # 49. Show a representative synthetic B inflow and the planned scenario grid.
    s = add_slide(prs, "Red B · hidrograma sintético y campaña SWMM",
                  "Ejemplo de entrada nodal para Chico Norte; los caudales provienen del INP y de formas sintéticas.", 49)
    s.shapes.add_picture(str(visuals["b_hydrograph"]), Inches(0.35), Inches(1.18), width=Inches(5.70))
    hydro = visuals["b_hydrograph_stats"]
    add_card(s, 6.28, 1.38, 3.18, 1.25, "Escenarios planificados", "975",
             "39 formas de hidrograma × 25 factores de intensidad.", fill=PALE_BLUE, accent=BLUE, value_size=23)
    add_card(s, 6.28, 2.88, 3.18, 1.25, "Grilla de factores", "0,2–5,0",
             "Paso de 0,2; se conserva el pico nodal de referencia y se escala por factor.",
             fill=PALE_BLUE, accent=BLUE, value_size=19)
    add_text(s, 0.60, 4.35, 8.85, 0.38,
             f"Ejemplo: nodo {hydro['node']} · pico base {fmt_decimal(hydro['base_peak_lps'], 2)} L/s · pico con factor {fmt_decimal(hydro['factor'], 1)}: {fmt_decimal(hydro['scenario_peak_lps'], 1)} L/s.",
             size=9.0, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, 0.58, 4.77, 8.88, 0.36,
             "Entrada sintética por nodo (no lluvia ni transformación lluvia–escorrentía). La campaña usa 521 nodos con caudal base positivo.",
             size=8.5, color=GRAY, align=PP_ALIGN.CENTER)

    # 50-51. Pair each eligible SWMM flood map with the previous promoted model.
    for page, case, case_label in [
        (50, visuals["a_maps"][0], "A · escenario inundado mediano"),
        (51, visuals["a_maps"][1], "A · escenario de mayor volumen total"),
    ]:
        stats = case["stats"]
        s = add_slide(prs, case_label,
                      f"{stats['shape_id']} · factor {fmt_decimal(stats['factor_mult'], 1)} · misma escala espacial en ambos mapas.", page)
        s.shapes.add_picture(str(case["swmm"]), Inches(0.30), Inches(1.04), width=Inches(4.58))
        s.shapes.add_picture(str(case["model"]), Inches(5.10), Inches(1.04), width=Inches(4.58))
        add_text(s, 0.48, 4.50, 4.22, 0.30,
                 f"SWMM · {fmt_int(stats['total_m3'])} m³ · {int(stats['n_flooded'])} nodos",
                 size=8.7, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, 5.28, 4.50, 4.22, 0.30,
                 f"XGBoost previo · {fmt_int(round(stats['model_total_m3']))} m³ · {stats['model_flooded_nodes']} nodos",
                 size=8.7, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        delta = stats["total_delta_m3"]
        add_text(s, 0.53, 4.88, 8.95, 0.27,
                 f"Diferencia del volumen total en este ejemplo: {delta:+,.0f} m³. Modelo promovido el 06-oct; comparación retrospectiva ilustrativa, no prueba independiente ni de transferencia.",
                 size=7.8, color=GRAY, align=PP_ALIGN.CENTER)

    # 52. XGBoost results.
    s = add_slide(prs, "XGBoost: 73,78 m\u00b3 de RMSE interno",
                  "50 ensayos Optuna · GroupKFold(5) por forma · semilla 42 · selecci\u00f3n solo con A", 52)
    s.shapes.add_picture(chart_bytes(study), Inches(0.34), Inches(1.10), width=Inches(6.10))
    add_card(s, 6.72, 1.24, 2.95, 1.18, "Mejor ensayo", "#13",
             "ID Optuna; numeraci\u00f3n desde cero.", fill=PALE_BLUE, accent=BLUE, value_size=23)
    add_card(s, 6.72, 2.64, 2.95, 1.18, "Criterio", "73,782 m\u00b3",
             "macro_network_rmse; en A equivale al RMSE de una red.", fill=PALE_BLUE, accent=BLUE, value_size=18)
    add_box(s, 6.72, 4.05, 2.95, 0.76, PALE_GOLD)
    add_text(s, 6.88, 4.13, 2.62, 0.60,
             "Candidato ajustado con A; a\u00fan no es el resultado de B.",
             size=10, color=NAVY, bold=True)
    add_text(s, 0.48, 4.80, 5.96, 0.21,
             "La l\u00ednea oscura muestra el mejor RMSE acumulado; es una estimaci\u00f3n de selecci\u00f3n interna.",
             size=8.5, color=GRAY)

    # 53. Model and B operational status.
    s = add_slide(prs, "Estado de modelos y simulaciones de la red B",
                  "Los conteos de B son operativos; todav\u00eda no son m\u00e9tricas de desempe\u00f1o.", 53)
    n_ok = len(b_progress.get("results", []))
    n_excluded = len(b_progress.get("failures", []))
    total = 39 * 25
    pending = max(total - n_ok - n_excluded, 0)
    add_box(s, 0.52, 1.25, 4.30, 2.78, PALE_BLUE)
    add_text(s, 0.76, 1.43, 3.82, 0.32, "XGBOOST · COMPLETADO", size=10.3, color=BLUE, bold=True)
    add_text(s, 0.76, 1.90, 3.80, 0.52, "50 / 50 ensayos", size=20, color=NAVY, bold=True)
    add_text(s, 0.76, 2.52, 3.78, 1.12,
             "Candidato ajustado con A. Se verificaron la recarga de los modelos y la procedencia de sus archivos.",
             size=10.2, color=BLACK, valign=MSO_ANCHOR.TOP)
    add_box(s, 5.18, 1.25, 4.30, 2.78, PALE_GOLD)
    add_text(s, 5.42, 1.43, 3.82, 0.32, "MLP · B\u00daSQUEDA ACTIVA", size=10.3, color=NAVY, bold=True)
    add_text(s, 5.42, 1.90, 3.80, 0.52, "50 ensayos previstos", size=20, color=NAVY, bold=True)
    add_text(s, 5.42, 2.52, 3.78, 1.12,
             "GroupKFold(5), 200 \u00e9pocas por fold. Falta ajustar el ensemble final de cinco semillas (42–46).",
             size=10.2, color=BLACK, valign=MSO_ANCHOR.TOP)
    add_text(s, 0.56, 4.23, 1.78, 0.32, "RED B · CORTE", size=9.8, color=BLUE, bold=True)
    add_text(s, 2.32, 4.20, 7.13, 0.39,
             f"{fmt_int(n_ok)} aptas · {fmt_int(n_excluded)} excluidas · {fmt_int(pending)} pendientes de {fmt_int(total)}",
             size=10.1, color=NAVY, bold=True)
    add_box(s, 0.58, 4.70, 8.88, 0.19, PALE_GRAY, square=True)
    progress = (n_ok + n_excluded) / total
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(4.70), Inches(8.88 * min(progress, 1)), Inches(0.19))
    bar.fill.solid(); bar.fill.fore_color.rgb = MID_BLUE; bar.line.color.rgb = MID_BLUE
    add_text(s, 0.58, 4.91, 8.88, 0.22,
             "Etiquetas de B a\u00fan cerradas para el scoring.", size=8.8, color=GRAY)

    # 54. Readiness of parity maps.
    s = add_slide(prs, "Mapas y gr\u00e1ficos de paridad: siguiente hito",
                  "La comparación con XGBoost previo es solo ilustrativa; la validación externa A→B sigue pendiente.", 54)
    steps = [
        (0.50, "1", "Terminar MLP", "Cerrar b\u00fasqueda y ajustar cinco semillas.", PALE_GOLD, NAVY),
        (2.82, "2", "Sellar", "Congelar modelos, variables y hashes de predicci\u00f3n.", PALE_BLUE, BLUE),
        (5.14, "3", "Evaluar B", "Verificar RPT, QA y cobertura; calcular m\u00e9tricas.", PALE_BLUE, BLUE),
        (7.46, "4", "Graficar", "Paridad por nodo/escenario y mapas espaciales.", PALE_GOLD, NAVY),
    ]
    for x, n, title, detail, fill, accent in steps:
        add_box(s, x, 1.48, 2.08, 1.70, fill)
        add_text(s, x + 0.13, 1.59, 0.34, 0.38, n, size=16, color=accent, bold=True)
        add_text(s, x + 0.48, 1.59, 1.46, 0.38, title, size=10.4, color=NAVY, bold=True)
        add_text(s, x + 0.15, 2.04, 1.78, 0.91, detail, size=8.9, color=BLACK, valign=MSO_ANCHOR.TOP)
    add_text(s, 0.76, 3.53, 8.48, 0.72,
             "La comparación con el modelo previo es retrospectiva; no reemplaza la paridad OOF ni la evaluación externa de B.",
             size=12, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, 0.80, 4.47, 8.72, 0.32,
             "Aún no reportamos % de confiabilidad: se calculará con predicciones selladas y etiquetas verificadas de B.",
             size=9.6, color=GRAY, align=PP_ALIGN.CENTER)

    # 55. Findings and limitations.
    s = add_slide(prs, "Qu\u00e9 podemos afirmar hoy y qu\u00e9 falta comprobar",
                  "La auditor\u00eda convirti\u00f3 una tabla hist\u00f3rica en un experimento trazable con etiquetas protegidas.", 55)
    for x, heading, bullets, fill, accent in [
        (0.48, "Evidencia disponible", [
            "Dataset A reconstruido y filtrado, con manifiestos y hashes.",
            "XGBoost seleccionado entre 50 ensayos; candidato verificado.",
            "Nueva red B en simulaci\u00f3n, con exclusiones reportadas."], PALE_BLUE, BLUE),
        (5.12, "Evidencia pendiente", [
            "Terminar MLP y ajustar el ensemble de cinco semillas.",
            "Sellar predicciones; evaluar B y generar mapas/paridad nuevos.",
            "Probar otra red antes de afirmar generalizaci\u00f3n amplia."], PALE_GOLD, NAVY),
    ]:
        add_box(s, x, 1.25, 4.39, 2.64, fill)
        add_text(s, x + 0.20, 1.40, 3.98, 0.36, heading, size=12.5, color=accent, bold=True)
        y = 1.91
        for bullet in bullets:
            add_text(s, x + 0.22, y, 3.95, 0.54, f"• {bullet}", size=9.7, color=BLACK, valign=MSO_ANCHOR.TOP)
            y += 0.59
    add_text(s, 0.61, 4.12, 9.0, 0.52,
             "Cautela hidr\u00e1ulica: continuidad ≤ 5%, extensiones de drenaje y tolerancias num\u00e9ricas son provisionales; falta justificarlas.",
             size=9.8, color=NAVY, bold=True)
    add_text(s, 0.61, 4.68, 9.0, 0.27,
             "Alcance: hidrogramas sint\u00e9ticos de caudal por nodo; lluvia–escorrent\u00eda sigue fuera del proyecto actual.",
             size=9.2, color=GRAY)

    # 56. Reproducibility anchors.
    s = add_slide(prs, "Trazabilidad y decisiones congeladas",
                  "Cada resultado queda asociado a una preparaci\u00f3n, un protocolo y artefactos verificables.", 56)
    rows = [
        ("ID de preparaci\u00f3n A", PREP_ID, "138.560 filas · 866 corridas · tabular_v4_21"),
        ("Manifiesto de A", "4bf0c50f61b2287…", "Fuentes, archivos por corrida y QA"),
        ("Prerregistro", "6d5b729e33bd1162…", "Sellado antes de las b\u00fasquedas"),
        ("Estudio XGBoost", "5b3ace9ece0223b8…", "50 ensayos · semilla 42 · GroupKFold(5)"),
        ("Protocolo B", str(b_progress.get("protocol_sha256", ""))[:16] + "…", "39 × 25 escenarios · staging sin etiquetas"),
    ]
    y = 1.24
    for i, (label, value, detail) in enumerate(rows):
        add_box(s, 0.48, y, 9.05, 0.63, PALE_GRAY if i % 2 == 0 else WHITE, square=True)
        add_text(s, 0.63, y + 0.08, 2.05, 0.40, label, size=9.2, color=GRAY, bold=True)
        add_text(s, 2.75, y + 0.08, 2.65, 0.40, value, size=9.8, color=NAVY, bold=True)
        add_text(s, 5.50, y + 0.08, 3.82, 0.40, detail, size=8.9, color=BLACK)
        y += 0.68
    add_text(s, 0.60, 4.80, 8.95, 0.22,
             "Rutas y hashes completos en los manifiestos y la bit\u00e1cora; las etiquetas de B no se guardan en SQLite.",
             size=8.6, color=GRAY)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(OUTPUT)
    print(f"slides={len(prs.slides)} | B aptas={n_ok} | B excluidas={n_excluded} | B pendientes={pending}")
    print(f"A aptas={quality['n_runs']} | filas={quality['n_rows']} | mejor XGBoost={study['best_value']:.6f} m3")
    print(f"prerregistro_sha={prereg.get('sha256', 'ver execution log')}")


if __name__ == "__main__":
    main()
