"""Pruebas del armado de la presentacion del 13 de octubre (avances)."""

import csv
import hashlib
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from scripts import armar_presentacion_avances_13_10 as deck  # noqa: E402

DATOS_REALES = (
    Path(__file__).resolve().parents[2] / "docs" / "auditoria" / "correcciones" / "datos"
)

CAMPOS = [
    "modelo", "variante", "protocolo", "que_mide", "f1_agregado", "recall_agregado",
    "precision_agregada", "auc_roc_agregado", "f1_media_por_fold_con_ceros",
    "folds_f1_indefinido", "rmse_vol_agregado_m3", "rmse_vol_media_por_fold_m3",
    "error_vol_total_pct", "nse_regresor_oracle",
]

# (modelo, variante, f1, auc, rmse, nse)
MODELOS_FALSOS = [
    ("baseline_forma_mas_cercana", "referencia", 0.9764, 0.9844, 172.86, 0.71),
    ("xgboost_default", "default", 0.9815, 0.9993, 112.80, 0.877),
    ("xgboost_tuned", "tuned", 0.9812, 0.9993, 69.64, 0.953),
    ("mlp_default", "default", 0.9783, 0.9992, 114.38, 0.878),
    ("mlp_tuned", "tuned", 0.9783, 0.9994, 101.12, 0.902),
    ("svm_default", "default", 0.8625, 0.9719, 222.46, 0.548),
    ("svm_tuned", "tuned", 0.8302, 0.9626, 214.87, 0.583),
]


def _escribir(ruta: Path, campos, filas):
    with open(ruta, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=campos)
        writer.writeheader()
        writer.writerows(filas)


def _fila(modelo, variante, protocolo, f1, auc, rmse, nse):
    return {
        "modelo": modelo, "variante": variante, "protocolo": protocolo, "que_mide": "x",
        "f1_agregado": f1, "recall_agregado": 0.99, "precision_agregada": 0.97,
        "auc_roc_agregado": auc, "f1_media_por_fold_con_ceros": 0.9,
        "folds_f1_indefinido": 0, "rmse_vol_agregado_m3": rmse,
        "rmse_vol_media_por_fold_m3": rmse, "error_vol_total_pct": -1.0,
        "nse_regresor_oracle": nse,
    }


def _datos_parte2_falsos(carpeta: Path) -> Path:
    """Tablas con las columnas y nombres de modelo REALES de la parte 2."""
    carpeta.mkdir(parents=True, exist_ok=True)
    filas = [_fila(m, v, "LeaveOneShapeOut", f1, auc, rmse, nse)
             for m, v, f1, auc, rmse, nse in MODELOS_FALSOS]
    filas.append(_fila("xgboost_tuned", "tuned", "LOSO", 0.9, 0.9, 50.0, 0.9))  # se ignora
    _escribir(carpeta / "comparacion_protocolos.csv", CAMPOS, filas)
    _escribir(carpeta / "matriz_confusion.csv", ["modelo", "protocolo", "tn", "fp", "fn", "tp"],
              [{"modelo": m, "protocolo": "LeaveOneShapeOut", "tn": 100000 + k, "fp": 1000 + k,
                "fn": 300 + k, "tp": 45000 + k} for k, (m, *_) in enumerate(MODELOS_FALSOS)])
    _escribir(carpeta / "banda_limitrofe.csv",
              ["modelo", "protocolo", "fn_total", "fn_en_banda_1_2m3", "fn_fuera_de_banda",
               "fn_con_vol_ge_100m3"],
              [{"modelo": m, "protocolo": "LeaveOneShapeOut", "fn_total": 300 + k,
                "fn_en_banda_1_2m3": 200 + k, "fn_fuera_de_banda": 100,
                "fn_con_vol_ge_100m3": 0} for k, (m, *_) in enumerate(MODELOS_FALSOS)])
    _escribir(carpeta / "resumen_afinacion.csv",
              ["family", "n_trials", "best_value", "duration_s", "fixed_params", "status"],
              [{"family": f, "n_trials": 3, "best_value": b, "duration_s": d,
                "fixed_params": '{"classifier": {"n_jobs": 1, "epochs": 200, "batch_size": 1024}}',
                "status": "estimacion_exploratoria_condicionada"}
               for f, b, d in (("svm", 200.0, 840.0), ("xgboost", 65.0, 756.0),
                               ("mlp", 77.0, 3100.0))])
    ensayos = {"svm": [220.0, 230.0, 210.0], "xgboost": [300.0, 100.0, 120.0],
               "mlp": [150.0, 90.0, 95.0]}
    _escribir(carpeta / "convergencia_afinacion.csv",
              ["family", "trial", "value", "duration_s", "best_so_far"],
              [{"family": f, "trial": t, "value": v, "duration_s": 1.0, "best_so_far": v}
               for f, valores in ensayos.items() for t, v in enumerate(valores)])
    return carpeta


def _prs_base():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(deck.SLIDE_W), Emu(deck.SLIDE_H)
    prs.slides.add_slide(prs.slide_layouts[6])
    return prs


def _textos(slide):
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
        if shape.has_table:
            out.extend(c.text for row in shape.table.rows for c in row.cells)
    return "\n".join(out)


def _titulo(slide):
    return next(s for s in slide.shapes if s.name == "TextBox 3").text_frame.text


def _pie(slide):
    return next(s for s in slide.shapes if s.name == "TextBox 5").text_frame.text


def _numero(slide):
    return next(s for s in slide.shapes if s.name == "TextBox 6").text_frame.text


def test_fmt_uses_english_number_format():
    assert deck.fmt(106526, 0) == "106,526"
    assert deck.fmt(9.1263, 2) == "9.13"
    assert deck.fmt(0.98170, 4) == "0.9817"
    assert deck.fmt(172.8642, 2) == "172.86"


def test_agregar_slide_resultados_agrega_cuatro_slides_con_los_nombres_reales(tmp_path):
    datos = _datos_parte2_falsos(tmp_path / "datos_parte2")
    prs = _prs_base()
    nuevas = deck.agregar_slide_resultados(prs, datos)
    assert len(nuevas) == 4 and len(prs.slides) == 5
    assert [_titulo(x) for x in nuevas] == [
        "XGBoost and MLP: results on unseen shapes",
        "Optuna tuning: 3 trials per family",
        "Confusion matrices: tuned XGBoost and MLP",
        "Conclusions and next steps",
    ]
    for slide in nuevas:
        assert not _titulo(slide).rstrip().endswith(".")
        assert len(slide.notes_slide.notes_text_frame.text) > 300
    assert [_numero(x) for x in nuevas] == ["2", "3", "4", "5"]

    resultados = nuevas[0]
    tabla = next(x for x in resultados.shapes if x.has_table).table
    celdas = [[c.text for c in row.cells] for row in tabla.rows]
    assert celdas[0] == ["Model", "RMSE", "F1", "AUC", "NSE"]
    assert [f[0] for f in celdas[1:]] == [
        "XGBoost tuned", "XGBoost default", "MLP tuned", "MLP default",
        "SVM tuned", "SVM default", "No-learning baseline"]
    assert celdas[1][1:] == ["69.64", "0.9812", "0.9993", "0.953"]
    graficos = [x.chart for x in resultados.shapes if x.has_chart]
    assert len(graficos) == 2  # RMSE y F1, por familia: por defecto vs afinado
    rmse = graficos[0].plots[0]
    assert list(rmse.categories) == ["XGBoost", "MLP", "Approximate SVM"]
    assert list(rmse.series[0].values) == [112.8, 114.38, 222.46]  # por defecto
    assert list(rmse.series[1].values) == [69.64, 101.12, 214.87]  # afinado
    assert "172.86" in _textos(resultados)  # baseline


def test_slide_de_afinacion_usa_el_minimo_acumulado(tmp_path):
    datos = _datos_parte2_falsos(tmp_path / "datos_parte2")
    _, afinacion, _, _ = deck.agregar_slide_resultados(_prs_base(), datos)
    grafico = next(x.chart for x in afinacion.shapes if x.has_chart)
    series = {s.name: list(s.values) for s in grafico.plots[0].series}
    assert series["XGBoost"] == [300.0, 100.0, 100.0]
    assert series["MLP"] == [150.0, 90.0, 90.0]
    assert series["Approximate SVM"] == [220.0, 220.0, 210.0]
    tabla = next(x for x in afinacion.shapes if x.has_table).table
    assert [c.text for c in tabla.rows[1].cells] == ["XGBoost", "65.0", "12.6"]
    assert "20 trials" not in _textos(afinacion)  # the number comes from the CSV (3)


def test_slide_de_matrices_usa_xgboost_y_mlp_afinados(tmp_path):
    datos = _datos_parte2_falsos(tmp_path / "datos_parte2")
    matrices = deck.agregar_slide_resultados(_prs_base(), datos)[2]
    texto = _textos(matrices)
    # xgboost_tuned es la fila 2 del CSV falso (k=2) y mlp_tuned la 4
    for valor in ("100,002", "1,002", "302", "45,002", "100,004", "1,004", "304", "45,004"):
        assert valor in texto
    assert "FN in the 1 m³ band: 202 of 302" in texto
    assert "FN in the 1 m³ band: 204 of 304" in texto
    assert "banda_limitrofe.csv" in texto


def test_conclusiones_leen_las_cifras_de_las_tablas(tmp_path):
    datos = _datos_parte2_falsos(tmp_path / "datos_parte2")
    conclusiones = deck.agregar_slide_resultados(_prs_base(), datos)[3]
    texto = _textos(conclusiones)
    assert "69.6 m³" in texto and "−38%" in texto and "101.1 m³" in texto
    assert "215 vs 173" in texto and "0.83 vs 0.976" in texto
    assert "200 epochs, batch 1024" in texto and "n_jobs=1" in texto
    assert "963 of 1000 runs" in texto  # without tables_provenance.json: default value


def test_agregar_slide_resultados_falla_si_falta_una_variante(tmp_path):
    datos = _datos_parte2_falsos(tmp_path / "datos_parte2")
    filas = [r for r in deck.cargar_csv(datos / "comparacion_protocolos.csv")
             if r["modelo"] != "mlp_tuned"]
    _escribir(datos / "comparacion_protocolos.csv", CAMPOS, filas)
    with pytest.raises(ValueError, match="mlp_tuned"):
        deck.agregar_slide_resultados(_prs_base(), datos)


def test_agregar_slide_resultados_falla_sin_filas_de_leave_one_shape_out(tmp_path):
    datos = tmp_path / "datos_parte2"
    datos.mkdir()
    _escribir(datos / "comparacion_protocolos.csv", CAMPOS,
              [_fila("xgboost_tuned", "tuned", "LOSO", 0.9, 0.9, 50.0, 0.9)])
    with pytest.raises(ValueError, match="LeaveOneShapeOut"):
        deck.agregar_slide_resultados(_prs_base(), datos)


def test_construir_sin_datos_parte2_agrega_seis_slides_y_no_toca_la_base(tmp_path):
    base = tmp_path / "base.pptx"
    _prs_base().save(str(base))
    huella = hashlib.sha256(base.read_bytes()).hexdigest()
    salida = tmp_path / "salida.pptx"

    prs = deck.construir(base, salida, DATOS_REALES, tmp_path / "no_existe")
    assert len(prs.slides) == 1 + 6
    assert hashlib.sha256(base.read_bytes()).hexdigest() == huella
    assert len(Presentation(str(salida)).slides) == 7

    # Idempotente: reconstruir da el mismo numero de slides, no acumula.
    prs2 = deck.construir(base, salida, DATOS_REALES, tmp_path / "no_existe")
    assert len(prs2.slides) == 7

    nuevas = list(prs2.slides)[1:]
    assert [_titulo(s) for s in nuevas] == [
        "Audit: what was found and what was fixed",
        "What each protocol measures (and why it matters)",
        "XGBoost confusion matrix on unseen shapes",
        "The flood label and its resolution",
        "Approximate SVM: architecture change",
        "How the fair comparison is run",
    ]
    for pos, slide in enumerate(nuevas, start=2):
        assert not _titulo(slide).rstrip().endswith(".")
        assert _numero(slide) == str(pos)
        assert _pie(slide).startswith(deck.FOOTER_TEXT)
        assert len(slide.notes_slide.notes_text_frame.text) > 300
    assert "results on the following slides" not in _pie(nuevas[-1])
    assert "to be measured" in _textos(nuevas[4])  # without resumen_afinacion.csv: expectation


def test_construir_con_datos_parte2_agrega_los_slides_de_resultados(tmp_path):
    base = tmp_path / "base.pptx"
    _prs_base().save(str(base))
    datos2 = _datos_parte2_falsos(tmp_path / "datos_parte2")
    prs = deck.construir(base, tmp_path / "salida.pptx", DATOS_REALES, datos2)
    assert len(prs.slides) == 1 + 6 + 4
    assert _pie(list(prs.slides)[6]).endswith("results on the following slides")
    # the SVM slide cites the measured cost (840 s = 14 min) from the tuning CSV
    assert "~14 min" in _textos(list(prs.slides)[5])
    assert "to be measured" not in _textos(list(prs.slides)[5])
    assert "3 Optuna trials per family" in _textos(list(prs.slides)[6])


def test_construir_se_niega_a_sobrescribir_la_base(tmp_path):
    base = tmp_path / "base.pptx"
    _prs_base().save(str(base))
    with pytest.raises(SystemExit):
        deck.construir(base, base, DATOS_REALES, tmp_path / "x")


def test_las_cifras_de_las_slides_de_auditoria_coinciden_con_los_csv():
    prs = _prs_base()
    cp = {(r["modelo"], r["protocolo"]): r
          for r in deck.cargar_csv(DATOS_REALES / "comparacion_protocolos.csv")}
    mc = {(r["modelo"], r["protocolo"]): r
          for r in deck.cargar_csv(DATOS_REALES / "matriz_confusion.csv")}
    protocolos = deck.slide_protocolos(prs, DATOS_REALES)
    matriz = deck.slide_matriz(prs, DATOS_REALES)

    claves = (
        ("baseline_interpolacion_factor", "LOSO"),
        ("baseline_forma_mas_cercana", "LeaveOneShapeOut"),
    )
    graficos = [s.chart for s in protocolos.shapes if s.has_chart]
    assert len(graficos) == 1
    series = graficos[0].plots[0].series
    assert list(series[0].values) == [
        round(float(cp[k]["rmse_vol_agregado_m3"]), 2) for k in claves]
    assert list(series[1].values) == [
        round(float(cp[("xgboost", p)]["rmse_vol_agregado_m3"]), 2)
        for p in ("LOSO", "LeaveOneShapeOut")]

    texto_m = _textos(matriz)
    m = mc[("xgboost", "LeaveOneShapeOut")]
    for campo in ("tn", "fp", "fn", "tp"):
        assert deck.fmt(m[campo], 0) in texto_m
    assert "333 of 380" in texto_m
    assert deck.fmt(cp[("xgboost", "LeaveOneShapeOut")]["f1_agregado"], 4) in texto_m


# --- English translation of the base deck ---------------------------------
def _prs_con_texto_es(tmp_path=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(deck.SLIDE_W), Emu(deck.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    caja = slide.shapes.add_textbox(0, 0, Emu(3000000), Emu(500000))
    caja.text_frame.text = "Presupuesto computacional: medir antes"
    tabla = slide.shapes.add_table(1, 2, 0, Emu(600000), Emu(3000000), Emu(400000)).table
    tabla.cell(0, 0).text = "Corridas SWMM"
    tabla.cell(0, 1).text = "1000"
    slide.notes_slide.notes_text_frame.text = "Fuente: https://example.org/x"
    return prs, slide


def test_translate_base_translates_frames_tables_and_notes():
    prs, slide = _prs_con_texto_es()
    traducciones = {
        "Presupuesto computacional: medir antes": "Computational budget: measure first",
        "Corridas SWMM": "SWMM runs",
        "Fuente: https://example.org/x": "Source: https://example.org/x",
    }
    contadores = deck.translate_base(prs, traducciones)
    assert contadores == {"frames": 2, "notes": 1, "alt": 0}
    texto = _textos(slide)
    assert "Computational budget: measure first" in texto and "SWMM runs" in texto
    assert "1000" in texto  # numbers are untouched
    assert slide.notes_slide.notes_text_frame.text == "Source: https://example.org/x"


def test_translate_base_fails_loudly_on_spanish_without_entry():
    prs, _ = _prs_con_texto_es()
    with pytest.raises(ValueError, match="Presupuesto computacional"):
        deck.translate_base(prs, {"Corridas SWMM": "SWMM runs"})


def test_translate_base_keeps_first_run_formatting_and_clears_the_rest():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    caja = slide.shapes.add_textbox(0, 0, Emu(3000000), Emu(500000))
    para = caja.text_frame.paragraphs[0]
    r1 = para.add_run()
    r1.text = "Morado: variables de "
    r1.font.bold = True
    r2 = para.add_run()
    r2.text = "entrada"
    deck.translate_base(prs, {"Morado: variables de entrada": "Purple: input variables"})
    runs = caja.text_frame.paragraphs[0].runs
    assert runs[0].text == "Purple: input variables" and runs[0].font.bold is True
    assert runs[1].text == ""


def test_the_real_base_deck_has_no_spanish_left_after_translating():
    base = deck.BASE
    if not base.exists():  # pragma: no cover
        pytest.skip("base deck not available")
    prs = Presentation(str(base))
    contadores = deck.translate_base(prs)  # raises if any Spanish paragraph lacks an entry
    assert contadores["frames"] > 150 and contadores["notes"] > 20
    assert not deck._parece_espanol(_textos(prs.slides[16]))  # slide 17 (was Spanish)
    assert next(
        s for s in prs.slides[36].shapes if s.name == "TextBox 3").text_frame.text == "Next steps"
